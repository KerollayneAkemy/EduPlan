from datetime import date
from html import escape
from io import BytesIO
import re


_BOLD = re.compile(r"\*\*(.+?)\*\*|(?<!\*)\*([^*]+)\*(?!\*)")
_HEADING = re.compile(r"^(#{1,3})\s+(.+)$")
_BULLET = re.compile(r"^[-*•●–]\s+(.+)$")


def _lines(content: str):
    for raw_line in content.splitlines():
        line = raw_line.strip()
        if not line:
            yield "blank", ""
            continue
        heading = _HEADING.match(line)
        if heading:
            yield f"heading{len(heading.group(1))}", heading.group(2).strip()
            continue
        bullet = _BULLET.match(line)
        if bullet:
            yield "indent", bullet.group(1).strip()
            continue
        yield "body", line


def _docx_runs(paragraph, text: str) -> None:
    cursor = 0
    for match in _BOLD.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor:match.start()])
        run = paragraph.add_run(match.group(1) or match.group(2))
        run.bold = True
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def _add_page_field(paragraph) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, value, end):
        run._r.append(element)


def _build_docx(material: dict) -> tuple[bytes, str, str]:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Inches, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Cm(3)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)
    section.header_distance = Inches(0.6)
    section.footer_distance = Inches(0.6)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.alignment = 3
    for name, size in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 12)):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(28, 74, 110) if name != "Heading 3" else RGBColor(45, 57, 69)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.first_line_indent = Inches(0)
        style.paragraph_format.alignment = 0
        style.paragraph_format.keep_with_next = True

    title = doc.add_paragraph(style="Title")
    title.alignment = 1
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(18)
    title.add_run(material["title"])
    title.runs[0].font.name = "Times New Roman"
    title.runs[0].font.size = Pt(16)
    title.runs[0].font.bold = True
    title.runs[0].font.color.rgb = RGBColor(24, 56, 82)

    details = doc.add_paragraph()
    details.paragraph_format.space_after = Pt(12)
    details.paragraph_format.first_line_indent = Inches(0)
    details.paragraph_format.alignment = 0
    detail_run = details.add_run(f"Disciplina: {material['subject']}\nTipo de material: {material['kind']}\nData de elaboração: {date.today().strftime('%d/%m/%Y')}")
    detail_run.font.name = "Times New Roman"
    detail_run.font.size = Pt(10)
    detail_run.font.color.rgb = RGBColor(0, 0, 0)

    for kind, text in _lines(material.get("content", "")):
        if kind == "blank":
            continue
        if kind.startswith("heading"):
            level = int(kind[-1])
            paragraph = doc.add_paragraph(style=f"Heading {level}")
            _docx_runs(paragraph, text)
        elif kind == "indent":
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(1.25)
            paragraph.paragraph_format.first_line_indent = Inches(0)
            _docx_runs(paragraph, text)
        else:
            paragraph = doc.add_paragraph()
            _docx_runs(paragraph, text)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.style = doc.styles["Caption"]
    footer.add_run("Página ")
    _add_page_field(footer)
    output = BytesIO()
    doc.save(output)
    return output.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"


def _build_pdf(material: dict) -> tuple[bytes, str, str]:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm, mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate

    output = BytesIO()
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="EduTitle", parent=styles["Title"], fontName="Times-Bold", fontSize=16, leading=20, alignment=TA_CENTER, textColor=colors.HexColor("#183852"), spaceAfter=16))
    styles.add(ParagraphStyle(name="EduMeta", parent=styles["Normal"], fontName="Times-Roman", fontSize=10, leading=14, textColor=colors.HexColor("#526577"), spaceAfter=12, alignment=TA_LEFT))
    styles.add(ParagraphStyle(name="EduBody", parent=styles["BodyText"], fontName="Times-Roman", fontSize=12, leading=18, alignment=TA_JUSTIFY, firstLineIndent=1.25 * cm, textColor=colors.black, spaceAfter=0, splitLongWords=True))
    styles.add(ParagraphStyle(name="EduIndented", parent=styles["EduBody"], leftIndent=1.25 * cm, firstLineIndent=0))
    styles.add(ParagraphStyle(name="EduH1", parent=styles["Heading1"], fontName="Times-Bold", fontSize=14, leading=19, textColor=colors.HexColor("#1C4A6E"), spaceBefore=14, spaceAfter=7, keepWithNext=True))
    styles.add(ParagraphStyle(name="EduH2", parent=styles["Heading2"], fontName="Times-Bold", fontSize=12, leading=18, textColor=colors.HexColor("#1C4A6E"), spaceBefore=11, spaceAfter=5, keepWithNext=True))
    styles.add(ParagraphStyle(name="EduH3", parent=styles["Heading3"], fontName="Times-Bold", fontSize=12, leading=18, textColor=colors.HexColor("#2D3945"), spaceBefore=9, spaceAfter=4, keepWithNext=True))

    def inline_markup(text: str) -> str:
        safe = escape(text)
        return _BOLD.sub(lambda match: f"<b>{match.group(1) or match.group(2)}</b>", safe)

    story = [
        Paragraph(escape(material["title"]), styles["EduTitle"]),
        Paragraph(f"Disciplina: {escape(material['subject'])}<br/>Tipo de material: {escape(material['kind'])}<br/>Data de elaboração: {date.today().strftime('%d/%m/%Y')}", styles["EduMeta"]),
    ]
    for kind, text in _lines(material.get("content", "")):
        if kind == "blank":
            continue
        elif kind.startswith("heading"):
            level = int(kind[-1])
            style = styles["EduH1"] if level == 1 else styles["EduH2"] if level == 2 else styles["EduH3"]
            story.append(Paragraph(inline_markup(text), style))
        elif kind == "indent":
            story.append(Paragraph(inline_markup(text), styles["EduIndented"]))
        else:
            story.append(Paragraph(inline_markup(text), styles["EduBody"]))

    def add_footer(canvas, document):
        canvas.saveState()
        canvas.setFont("Times-Roman", 10)
        canvas.setFillColor(colors.black)
        canvas.drawRightString(A4[0] - 20 * mm, 12 * mm, str(document.page))
        canvas.restoreState()

    SimpleDocTemplate(output, pagesize=A4, rightMargin=20 * mm, leftMargin=30 * mm, topMargin=30 * mm, bottomMargin=20 * mm, title=material["title"], author="EduPlan").build(story, onFirstPage=add_footer, onLaterPages=add_footer)
    return output.getvalue(), "application/pdf", "pdf"


def build_export(material: dict, format: str) -> tuple[bytes, str, str]:
    if format == "docx":
        return _build_docx(material)
    if format == "pdf":
        return _build_pdf(material)
    raise ValueError("Formato de exportação inválido. Use pdf ou docx.")
