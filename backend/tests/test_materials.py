from datetime import datetime
from io import BytesIO
import sys
from types import SimpleNamespace

import pytest
from docx import Document
from fastapi import HTTPException
from pypdf import PdfReader

from app.config import settings
from app.routes import materials as material_routes
from app.routes.materials import GenerateRequest, MaterialUpdate
from app.services.exports import build_export
from conftest import OTHER_ID, USER_ID


def test_material_list_is_scoped_to_authenticated_user(test_context):
    result = material_routes.list_materials(search=None, kind=None, subject=None, start=None, end=None, user={"id": USER_ID}, db=test_context, limit=100)
    assert [row["title"] for row in result] == ["Aula de Ciências"]


def test_material_shows_only_its_owners_syllabus_details(test_context):
    material = test_context.tables["materials"][0]
    material["syllabus_id"] = "55555555-5555-4555-8555-555555555555"
    test_context.tables["syllabi"].append({"id":material["syllabus_id"],"user_id":USER_ID,"subject":"Ciências","school_year":"6º ano","academic_year":2026,"semester":1,"file_name":"ciencias.docx"})
    result = material_routes.list_materials(search=None, kind=None, subject=None, start=None, end=None, user={"id":USER_ID}, db=test_context, limit=100)
    assert result[0]["syllabus"]["file_name"] == "ciencias.docx"


def test_search_filter_and_date_range(test_context):
    result = material_routes.list_materials(search="Ciências", kind=None, subject=None, start=datetime.fromisoformat("2026-09-01T00:00:00+00:00"), end=datetime.fromisoformat("2026-11-01T00:00:00+00:00"), user={"id":USER_ID}, db=test_context, limit=100)
    assert len(result) == 1


def test_cannot_read_or_edit_another_users_material(test_context):
    foreign_id = "44444444-4444-4444-8444-444444444444"
    with pytest.raises(HTTPException) as missing:
        material_routes.get_material(foreign_id, user={"id":USER_ID}, db=test_context)
    assert missing.value.status_code == 404
    with pytest.raises(HTTPException) as missing:
        material_routes.update_material(foreign_id, MaterialUpdate(title="Alterado",content="Alterado"), user={"id":USER_ID}, db=test_context)
    assert missing.value.status_code == 404


def test_edits_and_deletes_saved_material(test_context):
    material_id = "33333333-3333-4333-8333-333333333333"
    result = material_routes.update_material(material_id, MaterialUpdate(title="Plano revisado", content="Objetivo atualizado."), user={"id":USER_ID}, db=test_context)
    assert result["title"] == "Plano revisado"
    material_routes.delete_material(material_id, user={"id":USER_ID}, db=test_context)
    assert [row["id"] for row in test_context.tables["materials"]] == ["44444444-4444-4444-8444-444444444444"]


def test_exports_current_saved_content_as_pdf_and_docx(test_context):
    material = material_routes.get_material("33333333-3333-4333-8333-333333333333", user={"id":USER_ID}, db=test_context)
    pdf, pdf_type, _ = build_export(material, "pdf")
    docx, docx_type, _ = build_export(material, "docx")
    assert pdf_type == "application/pdf" and pdf.startswith(b"%PDF")
    assert docx_type.endswith("wordprocessingml.document") and docx.startswith(b"PK")


def test_exports_remove_markdown_markers_and_use_print_page_layout():
    material = {
        "title": "Atividade de leitura",
        "subject": "Língua Portuguesa",
        "kind": "atividade",
        "content": "## Leitura\n- Primeiro ponto\n* Segundo ponto\n**Texto em destaque**",
    }
    pdf, _, _ = build_export(material, "pdf")
    docx, _, _ = build_export(material, "docx")
    pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(pdf)).pages)
    word = Document(BytesIO(docx))
    word_text = "\n".join(paragraph.text for paragraph in word.paragraphs)
    for text in (pdf_text, word_text):
        assert "**" not in text
        assert "- Primeiro" not in text
        assert "* Segundo" not in text
    assert round(word.sections[0].left_margin / 360000) == 3
    assert round(word.sections[0].top_margin / 360000) == 3


def test_export_rejects_foreign_material(test_context):
    with pytest.raises(HTTPException) as missing:
        material_routes.export_material("44444444-4444-4444-8444-444444444444", "pdf", user={"id":USER_ID}, db=test_context)
    assert missing.value.status_code == 404


def test_generation_uses_mock_and_persists_result(test_context, monkeypatch):
    class Models:
        def generate_content(self, **kwargs):
            assert "Sugestões complementares" in kwargs["contents"]
            assert "Nível de ensino: Ensino Fundamental" in kwargs["contents"]
            assert kwargs["config"]["max_output_tokens"] == 4096
            return SimpleNamespace(text="## Objetivo\nRevisar o ciclo da água.")
    fake_types = SimpleNamespace(GenerateContentConfig=lambda **kwargs: kwargs)
    fake_genai = SimpleNamespace(Client=lambda **kwargs: SimpleNamespace(models=Models()), types=fake_types)
    monkeypatch.setitem(sys.modules, "google", SimpleNamespace(genai=fake_genai))
    monkeypatch.setitem(sys.modules, "google.genai", fake_genai)
    monkeypatch.setitem(sys.modules, "google.genai.types", fake_types)
    monkeypatch.setattr(settings, "gemini_api_key", "test-key")
    result = material_routes.generate(GenerateRequest(kind="atividade", subject="Ciências", topic="Ciclo da água", education_level="Ensino Fundamental", grade="6º ano"), user={"id":USER_ID}, db=test_context)
    assert result["content"].startswith("## Objetivo")
    assert any(row.get("content", "").startswith("## Objetivo") for row in test_context.tables["materials"])


def test_generation_returns_unavailable_without_gemini_key(test_context, monkeypatch):
    monkeypatch.setattr(settings, "gemini_api_key", "")
    with pytest.raises(HTTPException) as unavailable:
        material_routes.generate(GenerateRequest(kind="atividade",subject="Ciências",topic="Ciclo da água"), user={"id":USER_ID}, db=test_context)
    assert unavailable.value.status_code == 503


def test_generation_rejects_foreign_syllabus(test_context):
    test_context.tables["syllabi"].append({"id":"55555555-5555-4555-8555-555555555555","user_id":OTHER_ID,"extracted_text":"Privado"})
    with pytest.raises(HTTPException) as missing:
        material_routes.generate(GenerateRequest(kind="atividade",subject="Ciências",topic="Ciclo da água",syllabus_id="55555555-5555-4555-8555-555555555555"), user={"id":USER_ID}, db=test_context)
    assert missing.value.status_code == 404
